from datetime import date, datetime, timedelta
from io import BytesIO
import logging
from pathlib import Path
import re
from zipfile import ZipFile
import xml.etree.ElementTree as ET

from fastapi import APIRouter, BackgroundTasks, File, Form, HTTPException, UploadFile
from fastapi.responses import FileResponse
from openpyxl import load_workbook
from pydantic import ValidationError
from starlette.concurrency import run_in_threadpool

from app.database import clear_expired_timetables, get_connection
from app.schemas import (
    AiTimetableChange,
    AiTimetableEditRequest,
    AiTimetableEditResult,
    CourseInput,
    TimetableOut,
    TimetableRequest,
    VenueInput,
)
from app.services.ai_timetable import (
    AiTimetableConfigurationError,
    AiTimetableInterpretationError,
    AiTimetableTimeoutError,
    interpret_timetable_edit,
)
from app.services.excel_export import build_timetable_workbook
from app.services.timetable_edits import plan_ai_timetable_edits
from app.services.timetable import (
    generate_timetable_entries,
    generate_wizard_timetable_entries,
)

router = APIRouter(prefix="/allocation", tags=["allocation"])
logger = logging.getLogger(__name__)

_MAX_AI_CONTEXT_EXAMS = 250
_MAX_AI_EDIT_CHANGES = 200


@router.post("/generate", response_model=TimetableOut)
def generate_timetable(payload: TimetableRequest) -> TimetableOut:
    try:
        generated_entries = generate_timetable_entries(payload)
    except ValueError as error:
        raise HTTPException(status_code=400, detail=str(error))

    return _save_generated_timetable(payload, generated_entries)


@router.post("/generate/stored", response_model=TimetableOut)
def generate_timetable_from_stored_data() -> TimetableOut:
    payload = _stored_timetable_request()
    try:
        generated_entries = generate_timetable_entries(payload)
    except ValueError as error:
        raise HTTPException(status_code=400, detail=str(error))

    return _save_generated_timetable(payload, generated_entries, replace_existing=True)


@router.post("/generate/wizard", response_model=TimetableOut)
async def generate_timetable_from_wizard(
    title: str = Form("Smart Allocation Timetable"),
    start_date: str = Form(...),
    end_date: str = Form(...),
    semester: str = Form("First Semester"),
    time_slots: str = Form(...),
    lecturers: str = Form(""),
    venues: str = Form(""),
    courses: str = Form(""),
    programmes: str = Form(""),
    course_lecturers: str = Form(""),
    max_session_duration_minutes: int = Form(120),
    lecturer_files: list[UploadFile] | None = File(default=None),
    venue_files: list[UploadFile] | None = File(default=None),
    course_files: list[UploadFile] | None = File(default=None),
    programme_files: list[UploadFile] | None = File(default=None),
    course_lecturer_files: list[UploadFile] | None = File(default=None),
) -> TimetableOut:
    uploaded_data = await _extract_uploaded_documents(
        {
            "lecturers": lecturer_files,
            "venues": venue_files,
            "courses": course_files,
            "programmes": programme_files,
            "course_lecturers": course_lecturer_files,
        }
    )
    try:
        generated_entries = generate_wizard_timetable_entries(
            title=title,
            start_date=start_date,
            end_date=end_date,
            semester=semester,
            time_slots=time_slots,
            lecturers=_combine_source_text(lecturers, uploaded_data["lecturers"]),
            venues=_combine_source_text(venues, uploaded_data["venues"]),
            courses=_combine_source_text(courses, uploaded_data["courses"]),
            programmes=_combine_source_text(programmes, uploaded_data["programmes"]),
            course_lecturers=_combine_source_text(
                course_lecturers, uploaded_data["course_lecturers"]
            ),
            max_session_duration_minutes=max_session_duration_minutes,
        )
    except ValueError as error:
        raise HTTPException(status_code=400, detail=str(error))

    payload = TimetableRequest(
        title=title,
        exam_period=f"{start_date} to {end_date}",
        exam_dates=[],
        start_times=[],
        courses=[],
        venues=[],
        invigilators=[],
    )
    return _save_generated_timetable(payload, generated_entries, replace_existing=True)


def _save_generated_timetable(
    payload: TimetableRequest,
    generated_entries,
    *,
    replace_existing: bool = False,
) -> TimetableOut:
    with get_connection() as db:
        if replace_existing:
            db.execute("DELETE FROM timetables")

        cursor = db.execute(
            """
            INSERT INTO timetables (title, exam_period)
            VALUES (?, ?)
            RETURNING id, version, last_updated
            """,
            (payload.title, payload.exam_period),
        )
        timetable_row = cursor.fetchone()
        timetable_id = timetable_row["id"]
        saved_entries = []

        for entry in generated_entries:
            entry_cursor = db.execute(
                """
                INSERT INTO timetable_entries
                    (timetable_id, course_code, course_title, exam_date, start_time, end_time,
                     venue, student_level, student_count, invigilator_username)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                RETURNING id
                """,
                (
                    timetable_id,
                    entry.course_code,
                    entry.course_title,
                    entry.exam_date,
                    entry.start_time,
                    entry.end_time,
                    entry.venue,
                    entry.student_level,
                    entry.student_count,
                    entry.invigilator_username,
                ),
            )
            entry_id = entry_cursor.fetchone()["id"]
            db.execute(
                "INSERT INTO assignments (timetable_entry_id, invigilator_username) VALUES (?, ?)",
                (entry_id, entry.invigilator_username),
            )
            saved_entries.append(entry.model_copy(update={"id": entry_id}))

    timetable = TimetableOut(
        id=timetable_id,
        title=payload.title,
        exam_period=payload.exam_period,
        entries=saved_entries,
        version=timetable_row["version"],
        last_updated=str(timetable_row["last_updated"]),
    )
    build_timetable_workbook(
        {"id": timetable.id, "title": timetable.title, "exam_period": timetable.exam_period},
        [entry.model_dump() for entry in timetable.entries],
    )
    return timetable


@router.get("/timetables", response_model=list[TimetableOut])
def list_timetables() -> list[TimetableOut]:
    clear_expired_timetables()

    with get_connection() as db:
        timetables = db.execute(
            """
            SELECT id, title, exam_period, version, last_updated
            FROM timetables
            ORDER BY generated_at DESC, id DESC
            """
        ).fetchall()

        result = []
        for timetable in timetables:
            entries = db.execute(
                "SELECT * FROM timetable_entries WHERE timetable_id = ? ORDER BY exam_date, start_time",
                (timetable["id"],),
            ).fetchall()
            result.append(
                TimetableOut(
                    id=timetable["id"],
                    title=timetable["title"],
                    exam_period=timetable["exam_period"],
                    entries=[_entry_out(row) for row in entries],
                    version=timetable["version"],
                    last_updated=str(timetable["last_updated"]),
                )
            )
    return result


@router.get("/timetable/version")
def timetable_version() -> dict:
    """A small response clients can poll before downloading timetable data."""
    clear_expired_timetables()
    with get_connection() as db:
        timetable = db.execute(
            """
            SELECT id, version, last_updated
            FROM timetables
            ORDER BY generated_at DESC, id DESC
            LIMIT 1
            """
        ).fetchone()
    if timetable is None:
        return {"timetable_id": None, "version": 0, "last_updated": None}
    return {
        "timetable_id": timetable["id"],
        "version": timetable["version"],
        "last_updated": str(timetable["last_updated"]),
    }


@router.post("/timetable/ai-edit", response_model=AiTimetableEditResult)
async def ai_edit_timetable(
    payload: AiTimetableEditRequest,
    background_tasks: BackgroundTasks,
) -> AiTimetableEditResult:
    """Interpret an admin request with Gemini, then validate/write it locally."""
    state = await run_in_threadpool(_editable_timetable_state)
    if state is None:
        raise HTTPException(status_code=404, detail="There is no generated timetable to edit yet.")

    prompt = payload.prompt.strip()
    context = _ai_context(state, prompt)
    try:
        proposal = await interpret_timetable_edit(
            prompt=prompt,
            context=context,
        )
    except AiTimetableConfigurationError as error:
        raise HTTPException(status_code=503, detail=str(error))
    except AiTimetableTimeoutError as error:
        raise HTTPException(status_code=504, detail=str(error))
    except AiTimetableInterpretationError as error:
        raise HTTPException(status_code=502, detail=str(error))
    except Exception as error:
        logger.exception("Unexpected Gemini failure while interpreting a timetable edit")
        raise HTTPException(
            status_code=502,
            detail="The AI timetable service failed unexpectedly. Please try again.",
        ) from error

    raw_changes = proposal.get("changes", [])
    if len(raw_changes) > _MAX_AI_EDIT_CHANGES:
        raise HTTPException(
            status_code=422,
            detail=(
                f"The AI proposed {len(raw_changes)} changes. Refine the request to no more "
                f"than {_MAX_AI_EDIT_CHANGES} exams at a time."
            ),
        )

    changes: list[AiTimetableChange] = []
    rejected: list[dict] = []
    for raw_change in raw_changes:
        try:
            changes.append(AiTimetableChange.model_validate(raw_change))
        except ValidationError:
            rejected.append(
                {
                    "operation": "unknown",
                    "reason": "The AI returned an incomplete change that could not be validated.",
                }
            )

    # Do not apply a proposal to a timetable that changed while Gemini was
    # interpreting the request. The admin can safely submit it again.
    latest_state = await run_in_threadpool(_editable_timetable_state)
    if latest_state is None or (
        latest_state["id"] != state["id"]
        or latest_state["version"] != state["version"]
    ):
        raise HTTPException(
            status_code=409,
            detail="The timetable changed while the AI was preparing this edit. Please try again.",
        )

    final_entries, applied, plan_rejected = await run_in_threadpool(
        plan_ai_timetable_edits,
        changes=changes,
        entries=state["entries"],
        venues=state["venues"],
        invigilators=state["invigilators"],
        course_catalogue=state["course_catalogue"],
    )
    rejected.extend(plan_rejected)
    unresolved = [str(item) for item in proposal.get("unresolved", []) if str(item).strip()]

    if not applied:
        return AiTimetableEditResult(
            applied=[],
            rejected=rejected,
            unresolved=unresolved,
            version=state["version"],
            last_updated=state["last_updated"],
        )

    version, last_updated = await run_in_threadpool(
        _save_ai_timetable_edits,
        state=state,
        final_entries=final_entries,
    )
    # Excel generation can be slow for large timetables, so it must not keep
    # this API request open after the validated database edit is committed.
    background_tasks.add_task(_export_current_timetable, state["id"])
    return AiTimetableEditResult(
        applied=applied,
        rejected=rejected,
        unresolved=unresolved,
        version=version,
        last_updated=last_updated,
    )


@router.delete("/timetables")
def clear_timetables() -> dict[str, str]:
    with get_connection() as db:
        db.execute("DELETE FROM timetables")
    return {"message": "Generated timetable cleared."}


@router.get("/timetable/{timetable_id}/excel")
def download_timetable_excel(timetable_id: int) -> FileResponse:
    clear_expired_timetables()

    with get_connection() as db:
        timetable = db.execute(
            "SELECT id, title, exam_period FROM timetables WHERE id = ?",
            (timetable_id,),
        ).fetchone()
        if timetable is None:
            raise HTTPException(status_code=404, detail="Timetable not found.")

        entries = db.execute(
            "SELECT * FROM timetable_entries WHERE timetable_id = ? ORDER BY exam_date, start_time",
            (timetable_id,),
        ).fetchall()

    path = build_timetable_workbook(dict(timetable), [dict(row) for row in entries])
    return FileResponse(
        path,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=path.name,
    )


def _entry_out(row) -> dict:
    return {
        "id": row["id"],
        "course_code": row["course_code"],
        "course_title": row["course_title"],
        "exam_date": str(row["exam_date"]),
        "start_time": str(row["start_time"])[:5],
        "end_time": str(row["end_time"])[:5],
        "venue": row["venue"],
        "student_level": row["student_level"],
        "student_count": row["student_count"],
        "invigilator_username": row["invigilator_username"],
    }


def _editable_timetable_state() -> dict | None:
    """Load the latest timetable and the catalogue needed to validate edits."""
    clear_expired_timetables()
    with get_connection() as db:
        timetable = db.execute(
            """
            SELECT id, title, exam_period, version, last_updated
            FROM timetables
            ORDER BY generated_at DESC, id DESC
            LIMIT 1
            """
        ).fetchone()
        if timetable is None:
            return None
        rows = db.execute(
            "SELECT * FROM timetable_entries WHERE timetable_id = ? ORDER BY id",
            (timetable["id"],),
        ).fetchall()
        venue_rows = db.execute(
            "SELECT name, capacity FROM allocation_venues ORDER BY id"
        ).fetchall()
        invigilator_rows = db.execute(
            """
            SELECT username FROM allocation_invigilators
            UNION
            SELECT username FROM users WHERE role = 'invigilator'
            """
        ).fetchall()
        course_rows = db.execute(
            """
            SELECT course_code, course_title, student_level, student_count, duration_minutes
            FROM allocation_courses
            """
        ).fetchall()

    entries = [_entry_out(row) for row in rows]
    venue_capacities = {row["name"].lower(): (row["name"], row["capacity"]) for row in venue_rows}
    # Wizard uploads are not persisted as venue records.  Existing used venues
    # remain editable at their already-scheduled capacity, while new unknown
    # rooms are rejected until their capacity has been configured.
    for entry in entries:
        key = entry["venue"].lower()
        current = venue_capacities.get(key)
        if current is None:
            venue_capacities[key] = (entry["venue"], entry["student_count"])
        elif entry["student_count"] > current[1]:
            venue_capacities[key] = (current[0], entry["student_count"])
    venues = [VenueInput(name=name, capacity=capacity) for name, capacity in venue_capacities.values()]

    invigilators = {row["username"] for row in invigilator_rows}
    invigilators.update(entry["invigilator_username"] for entry in entries)
    course_catalogue = {
        row["course_code"].lower(): dict(row) for row in course_rows
    }
    for entry in entries:
        key = entry["course_code"].lower()
        course_catalogue.setdefault(
            key,
            {
                "course_code": entry["course_code"],
                "course_title": entry["course_title"],
                "student_level": entry["student_level"],
                "student_count": entry["student_count"],
                "duration_minutes": _entry_duration_minutes(entry),
            },
        )

    return {
        "id": timetable["id"],
        "title": timetable["title"],
        "exam_period": timetable["exam_period"],
        "version": timetable["version"],
        "last_updated": str(timetable["last_updated"]),
        "entries": entries,
        "venues": venues,
        "invigilators": sorted(invigilators),
        "course_catalogue": course_catalogue,
    }


def _ai_context(state: dict, prompt: str) -> dict:
    entries = _entries_relevant_to_prompt(state["entries"], prompt)
    if entries is None:
        raise HTTPException(
            status_code=422,
            detail=(
                "This timetable is too large to send in full to the AI. Include a course "
                "code, course title, venue, or invigilator name in the request."
            ),
        )
    if len(entries) > _MAX_AI_CONTEXT_EXAMS:
        raise HTTPException(
            status_code=422,
            detail=(
                f"Your request matches {len(entries)} exams. Refine it to at most "
                f"{_MAX_AI_CONTEXT_EXAMS} exams so the AI can respond reliably."
            ),
        )

    selected_course_codes = {entry["course_code"].lower() for entry in entries}
    return {
        "timetable_id": state["id"],
        "version": state["version"],
        "title": state["title"],
        "exam_period": state["exam_period"],
        "exams": entries,
        "known_venues": [
            {"name": venue.name, "capacity": venue.capacity}
            for venue in state["venues"]
        ],
        "known_invigilators": state["invigilators"],
        "known_courses": [
            course
            for code, course in state["course_catalogue"].items()
            if code in selected_course_codes
        ],
    }


def _entries_relevant_to_prompt(entries: list[dict], prompt: str) -> list[dict] | None:
    if len(entries) <= _MAX_AI_CONTEXT_EXAMS:
        return entries

    ignored_terms = {
        "all", "and", "change", "date", "exam", "exams", "from", "move", "next",
        "please", "schedule", "the", "this", "time", "timetable", "to", "week", "with",
    }
    terms = {
        term
        for term in re.findall(r"[a-z0-9]+", prompt.casefold())
        if len(term) >= 3 and term not in ignored_terms
    }
    if not terms:
        return None

    relevant = []
    for entry in entries:
        searchable = " ".join(
            str(entry[field]).casefold()
            for field in ("course_code", "course_title", "venue", "invigilator_username")
        )
        if any(term in searchable for term in terms):
            relevant.append(entry)
    return relevant or None


def _entry_duration_minutes(entry: dict) -> int:
    start = datetime.strptime(entry["start_time"][:5], "%H:%M")
    end = datetime.strptime(entry["end_time"][:5], "%H:%M")
    return int((end - start).total_seconds() // 60)


def _save_ai_timetable_edits(*, state: dict, final_entries: list[dict]) -> tuple[int, str]:
    original_by_id = {entry["id"]: entry for entry in state["entries"]}
    with get_connection() as db:
        # Claim the exact version in one statement before modifying entries.
        # A separate SELECT followed by a later UPDATE leaves a race where two
        # requests can both read the same version and overwrite one another.
        updated = db.execute(
            """
            UPDATE timetables
            SET version = version + 1, last_updated = CURRENT_TIMESTAMP::text
            WHERE id = ? AND version = ?
            RETURNING version, last_updated
            """,
            (state["id"], state["version"]),
        ).fetchone()
        if updated is None:
            raise HTTPException(
                status_code=409,
                detail="The timetable changed before the edit could be saved. Please try again.",
            )

        for entry in final_entries:
            if entry["id"] < 0:
                created = db.execute(
                    """
                    INSERT INTO timetable_entries
                        (timetable_id, course_code, course_title, exam_date, start_time, end_time,
                         venue, student_level, student_count, invigilator_username)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    RETURNING id
                    """,
                    (
                        state["id"], entry["course_code"], entry["course_title"],
                        entry["exam_date"], entry["start_time"], entry["end_time"],
                        entry["venue"], entry["student_level"], entry["student_count"],
                        entry["invigilator_username"],
                    ),
                ).fetchone()
                db.execute(
                    "INSERT INTO assignments (timetable_entry_id, invigilator_username) VALUES (?, ?)",
                    (created["id"], entry["invigilator_username"]),
                )
                continue
            if entry == original_by_id[entry["id"]]:
                continue
            db.execute(
                """
                UPDATE timetable_entries
                SET course_code = ?, course_title = ?, exam_date = ?, start_time = ?, end_time = ?,
                    venue = ?, student_level = ?, student_count = ?, invigilator_username = ?
                WHERE id = ? AND timetable_id = ?
                """,
                (
                    entry["course_code"], entry["course_title"], entry["exam_date"],
                    entry["start_time"], entry["end_time"], entry["venue"],
                    entry["student_level"], entry["student_count"], entry["invigilator_username"],
                    entry["id"], state["id"],
                ),
            )
            db.execute(
                "UPDATE assignments SET invigilator_username = ? WHERE timetable_entry_id = ?",
                (entry["invigilator_username"], entry["id"]),
            )

    return updated["version"], str(updated["last_updated"])


def _export_current_timetable(timetable_id: int) -> None:
    with get_connection() as db:
        timetable = db.execute(
            "SELECT id, title, exam_period FROM timetables WHERE id = ?",
            (timetable_id,),
        ).fetchone()
        entries = db.execute(
            "SELECT * FROM timetable_entries WHERE timetable_id = ? ORDER BY exam_date, start_time",
            (timetable_id,),
        ).fetchall()
    if timetable is not None:
        build_timetable_workbook(dict(timetable), [dict(entry) for entry in entries])


async def _extract_uploaded_documents(
    files_by_type: dict[str, list[UploadFile] | None],
) -> dict[str, str]:
    extracted = {
        "lecturers": [],
        "venues": [],
        "courses": [],
        "programmes": [],
        "course_lecturers": [],
    }
    for source_type, files in files_by_type.items():
        for upload in files or []:
            filename = upload.filename or "uploaded-file"
            content = await upload.read()
            if content:
                extracted[source_type].append(_extract_document_text(filename, content))
    return {source_type: "\n\n".join(notes) for source_type, notes in extracted.items()}


def _combine_source_text(typed_text: str, uploaded_text: str) -> str:
    return "\n".join(value for value in (typed_text.strip(), uploaded_text.strip()) if value)


def _extract_document_text(filename: str, content: bytes) -> str:
    extension = Path(filename).suffix.lower()
    try:
        if extension in {".xlsx", ".xlsm"}:
            text = _extract_excel_text(content)
        elif extension == ".docx":
            text = _extract_docx_text(content)
        elif extension == ".pdf":
            text = _extract_pdf_text(content)
        elif extension in {".txt", ".csv"}:
            text = content.decode("utf-8", errors="ignore")
        else:
            text = ""
    except Exception as error:
        text = f"Could not extract text from this file: {error}"

    # Keep all extractable records for the deterministic local scheduler.
    extracted_text = text.strip()
    if not extracted_text:
        extracted_text = "Uploaded file accepted, but no readable text could be extracted."
    return f"File: {filename}\n{extracted_text}"


def _extract_excel_text(content: bytes) -> str:
    workbook = load_workbook(BytesIO(content), data_only=True, read_only=True)
    lines = []
    for sheet in workbook.worksheets:
        lines.append(f"Sheet: {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            # Keep empty cells. Dropping them shifts later values into the
            # wrong columns (for example, a capacity can become a venue name).
            values = ["" if value is None else str(value).strip() for value in row]
            while values and not values[-1]:
                values.pop()
            if any(values):
                lines.append(" | ".join(values))
    return "\n".join(lines)


def _extract_docx_text(content: bytes) -> str:
    with ZipFile(BytesIO(content)) as archive:
        try:
            document = archive.read("word/document.xml")
        except KeyError:
            return ""
    root = ET.fromstring(document)
    namespace = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    paragraphs = []
    for paragraph in root.findall("./w:body/w:p", namespace):
        text = "".join(node.text or "" for node in paragraph.findall(".//w:t", namespace))
        if text.strip():
            paragraphs.append(text.strip())
    for table_row in root.findall(".//w:tr", namespace):
        cells = []
        for cell in table_row.findall(".//w:tc", namespace):
            cell_text = "".join(
                node.text or "" for node in cell.findall(".//w:t", namespace)
            ).strip()
            cells.append(cell_text)
        while cells and not cells[-1]:
            cells.pop()
        if any(cells):
            paragraphs.append(" | ".join(cells))
    return "\n".join(paragraphs)


def _extract_pdf_text(content: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError:
        return "PDF uploaded. Install pypdf on the backend to extract PDF text."

    reader = PdfReader(BytesIO(content))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _stored_timetable_request() -> TimetableRequest:
    with get_connection() as db:
        settings = db.execute("SELECT * FROM allocation_settings WHERE id = 1").fetchone()
        courses = db.execute(
            """
            SELECT course_code, course_title, student_level, student_count, duration_minutes
            FROM allocation_courses
            ORDER BY id
            """
        ).fetchall()
        venues = db.execute(
            "SELECT name, capacity FROM allocation_venues ORDER BY id"
        ).fetchall()
        invigilators = db.execute(
            """
            SELECT username, MIN(sort_order) AS sort_order
            FROM (
                SELECT username, id AS sort_order
                FROM allocation_invigilators
                UNION ALL
                SELECT username, 1000000 + id AS sort_order
                FROM users
                WHERE role = 'invigilator'
            )
            GROUP BY lower(username)
            ORDER BY sort_order
            """
        ).fetchall()

    if settings is None:
        raise HTTPException(status_code=400, detail="Allocation settings are not configured.")
    if date.today() > date.fromisoformat(settings["exam_end_date"]):
        clear_expired_timetables()
        raise HTTPException(
            status_code=400,
            detail="The exam period is over, so the timetable has been cleared.",
        )

    return TimetableRequest(
        title=settings["title"],
        exam_period=settings["exam_period"],
        exam_dates=_weekly_exam_dates(
            settings["exam_start_date"],
            settings["exam_end_date"],
            max(len(courses), len(invigilators)),
        ),
        start_times=[
            item.strip()
            for item in settings["start_times"].split(",")
            if item.strip()
        ],
        courses=[
            CourseInput(
                course_code=row["course_code"],
                course_title=row["course_title"],
                student_level=row["student_level"],
                student_count=row["student_count"],
                duration_minutes=row["duration_minutes"],
            )
            for row in courses
        ],
        venues=[
            VenueInput(name=row["name"], capacity=row["capacity"])
            for row in venues
        ],
        invigilators=[row["username"] for row in invigilators],
    )


def _weekly_exam_dates(start_date: str, end_date: str, needed_count: int) -> list[str]:
    start = datetime.strptime(start_date, "%Y-%m-%d").date()
    end = datetime.strptime(end_date, "%Y-%m-%d").date()
    dates: list[date] = []
    current = start

    while current <= end and len(dates) < needed_count:
        dates.append(current)
        current += timedelta(days=7)

    if not dates:
        return []

    while len(dates) < needed_count:
        dates.append(dates[-1])

    return [item.isoformat() for item in dates]
