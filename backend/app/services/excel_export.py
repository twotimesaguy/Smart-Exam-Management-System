from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins


_COLUMN_HEADERS = [
    "S/N", "Course Code", "Course Title", "Programme / Level", "Duration",
    "No. Reg.", "Venue", "Lead Invigilator",
]
_COLUMN_WIDTHS = [7, 16, 42, 21, 18, 11, 30, 28]
_TITLE_FILL = PatternFill("solid", fgColor="17365D")
_SESSION_FILL = PatternFill("solid", fgColor="1F4E78")
_HEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
_WHITE_FONT = Font(color="FFFFFF", bold=True)
_THIN_BLUE_BORDER = Border(
    left=Side(style="thin", color="8EA9DB"),
    right=Side(style="thin", color="8EA9DB"),
    top=Side(style="thin", color="8EA9DB"),
    bottom=Side(style="thin", color="8EA9DB"),
)


def build_timetable_workbook(timetable: dict, entries: list[dict]) -> Path:
    """Create a printable timetable grouped by exam date and session."""
    export_dir = Path(__file__).resolve().parents[2] / "exports"
    export_dir.mkdir(parents=True, exist_ok=True)
    path = export_dir / f"timetable_{timetable['id']}.xlsx"

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Exam Timetable"
    sheet.sheet_view.showGridLines = False
    _configure_page(sheet)
    _write_document_heading(sheet, timetable, entries)

    current_row = 5
    for (exam_date, start_time), session_entries in _group_entries(entries):
        _write_session_heading(sheet, current_row, exam_date, start_time)
        _write_column_headings(sheet, current_row + 1)
        current_row += 2
        for number, entry in enumerate(session_entries, start=1):
            _write_exam_row(sheet, current_row, number, entry)
            current_row += 1
        current_row += 1

    if not entries:
        sheet.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=8)
        cell = sheet.cell(current_row, 1, "No exam entries were generated.")
        cell.alignment = Alignment(horizontal="center")

    workbook.save(path)
    return path


def _configure_page(sheet) -> None:
    for index, width in enumerate(_COLUMN_WIDTHS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A5"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_margins = PageMargins(left=0.2, right=0.2, top=0.4, bottom=0.4)


def _write_document_heading(sheet, timetable: dict, entries: list[dict]) -> None:
    heading_rows = [
        str(timetable["title"]).upper(),
        "EXAMINATION TIMETABLE",
        f"EXAM PERIOD: {timetable['exam_period']}  |  EXAM DURATION: {_duration_label(entries)}",
    ]
    for row, value in enumerate(heading_rows, start=1):
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
        cell = sheet.cell(row, 1, value)
        cell.fill = _TITLE_FILL
        cell.font = Font(color="FFFFFF", bold=True, size=15 if row == 1 else 11)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        sheet.row_dimensions[row].height = 27 if row == 1 else 21


def _group_entries(entries: list[dict]) -> list[tuple[tuple[str, str], list[dict]]]:
    groups: dict[tuple[str, str], list[dict]] = defaultdict(list)
    for entry in entries:
        groups[(str(entry["exam_date"]), str(entry["start_time"]))].append(entry)
    return [
        (slot, sorted(group, key=lambda entry: (str(entry["venue"]), str(entry["course_code"]))))
        for slot, group in sorted(groups.items(), key=lambda item: (_sortable_date(item[0][0]), item[0][1]))
    ]


def _write_session_heading(sheet, row: int, exam_date: str, start_time: str) -> None:
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
    cell = sheet.cell(row, 1, f"{_display_date(exam_date)}: {_session_name(start_time)} SESSION")
    cell.fill = _SESSION_FILL
    cell.font = _WHITE_FONT
    cell.alignment = Alignment(horizontal="left", vertical="center")
    sheet.row_dimensions[row].height = 23


def _write_column_headings(sheet, row: int) -> None:
    for column, value in enumerate(_COLUMN_HEADERS, start=1):
        cell = sheet.cell(row, column, value)
        cell.fill = _HEADER_FILL
        cell.font = Font(bold=True, color="17365D")
        cell.border = _THIN_BLUE_BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[row].height = 30


def _write_exam_row(sheet, row: int, number: int, entry: dict) -> None:
    values = [
        number, entry["course_code"], entry["course_title"], entry["student_level"],
        f"{entry['start_time']} - {entry['end_time']}", entry["student_count"],
        entry["venue"], entry["invigilator_username"],
    ]
    for column, value in enumerate(values, start=1):
        cell = sheet.cell(row, column, value)
        cell.border = _THIN_BLUE_BORDER
        cell.alignment = Alignment(
            horizontal="left" if column in (3, 4, 7, 8) else "center",
            vertical="center", wrap_text=True,
        )
    sheet.row_dimensions[row].height = 34


def _session_name(start_time: str) -> str:
    try:
        hour = datetime.strptime(start_time[:5], "%H:%M").hour
    except ValueError:
        return "EXAM"
    if hour < 12:
        return "MORNING"
    if hour < 16:
        return "AFTERNOON"
    return "EVENING"


def _display_date(exam_date: str) -> str:
    try:
        return datetime.strptime(exam_date[:10], "%Y-%m-%d").strftime("%A %d %B %Y").upper()
    except ValueError:
        return exam_date.upper()


def _sortable_date(exam_date: str) -> datetime:
    try:
        return datetime.strptime(exam_date[:10], "%Y-%m-%d")
    except ValueError:
        return datetime.max


def _duration_label(entries: list[dict]) -> str:
    durations = {
        _duration_minutes(entry["start_time"], entry["end_time"])
        for entry in entries
        if entry.get("start_time") and entry.get("end_time")
    }
    if len(durations) != 1:
        return "VARIES"
    minutes = durations.pop()
    hours, remaining_minutes = divmod(minutes, 60)
    if hours and remaining_minutes:
        return f"{hours} HOURS {remaining_minutes} MINUTES"
    if hours:
        return f"{hours} HOURS"
    return f"{remaining_minutes} MINUTES"


def _duration_minutes(start_time: str, end_time: str) -> int:
    start = datetime.strptime(str(start_time)[:5], "%H:%M")
    end = datetime.strptime(str(end_time)[:5], "%H:%M")
    return int((end - start).total_seconds() // 60)
